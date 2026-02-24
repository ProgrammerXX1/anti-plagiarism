#!/usr/bin/env python3
"""
Генератор Word-документа: полное описание системы Anti-Plagiarism.
"""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

# ── Стили ──
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(4)

for level in range(1, 5):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Times New Roman'
    hs.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        hs.font.size = Pt(18)
        hs.font.bold = True
    elif level == 2:
        hs.font.size = Pt(16)
        hs.font.bold = True
    elif level == 3:
        hs.font.size = Pt(14)
        hs.font.bold = True
    else:
        hs.font.size = Pt(13)
        hs.font.bold = True

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # header row
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.name = 'Times New Roman'
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'D9E2F3')
        cell._tc.get_or_add_tcPr().append(shading)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri+1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.name = 'Times New Roman'
    if col_widths:
        for ri, row_obj in enumerate(table.rows):
            for ci, w in enumerate(col_widths):
                row_obj.cells[ci].width = Cm(w)
    return table

def add_code_block(doc, code, lang=''):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    # light gray background
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:fill'), 'F2F2F2')
    run._r.get_or_add_rPr().append(shading)

# ═══════════════════════════════════════════════════
# ТИТУЛЬНАЯ СТРАНИЦА
# ═══════════════════════════════════════════════════

for _ in range(6):
    doc.add_paragraph()

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run('СИСТЕМА АНТИПЛАГИАТА')
run.bold = True
run.font.size = Pt(26)
run.font.name = 'Times New Roman'

subtitle_p = doc.add_paragraph()
subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle_p.add_run('Anti-Plagiarism Pipeline')
run.font.size = Pt(18)
run.font.name = 'Times New Roman'
run.font.color.rgb = RGBColor(80, 80, 80)

doc.add_paragraph()

desc_p = doc.add_paragraph()
desc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = desc_p.add_run('Техническая документация\nАрхитектура, структура модулей, потоки данных,\nформаты хранения и алгоритмы работы')
run.font.size = Pt(14)
run.font.name = 'Times New Roman'

for _ in range(6):
    doc.add_paragraph()

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = date_p.add_run(f'Дата: {datetime.date.today().strftime("%d.%m.%Y")}')
run.font.size = Pt(12)
run.font.name = 'Times New Roman'

doc.add_page_break()

# ═══════════════════════════════════════════════════
# СОДЕРЖАНИЕ (вручную)
# ═══════════════════════════════════════════════════

doc.add_heading('СОДЕРЖАНИЕ', level=1)

toc_items = [
    ('1.', 'Общее описание системы'),
    ('2.', 'Высокоуровневая архитектура'),
    ('3.', 'Структура проекта (дерево файлов)'),
    ('4.', 'Модуль 1: Text — Нормализация и токенизация'),
    ('5.', 'Модуль 2: Builder — Построение сегментов'),
    ('6.', 'Модуль 3: Index_IO — Хранение на диске'),
    ('7.', 'Модуль 4: Search — Поиск плагиата'),
    ('8.', 'Модуль 5: Service — HTTP-сервис и воркеры'),
    ('9.', 'Система кэширования (Hot Cache)'),
    ('10.', 'PostgreSQL — метаданные и очередь задач'),
    ('11.', 'Пайплайн индексации (Pipeline)'),
    ('12.', 'Пайплайн поиска (Search Pipeline)'),
    ('13.', 'Формат хранения данных на диске'),
    ('14.', 'Подготовка корпуса и инджест скрипты'),
    ('15.', 'Лемматизация — Apertium Kazakh FST'),
    ('16.', 'Docker и деплой'),
    ('17.', 'API роуты'),
    ('18.', 'Ресурсы и производительность'),
]

for num, title in toc_items:
    p = doc.add_paragraph()
    run = p.add_run(f'{num} {title}')
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 1. ОБЩЕЕ ОПИСАНИЕ СИСТЕМЫ
# ═══════════════════════════════════════════════════

doc.add_heading('1. Общее описание системы', level=1)

doc.add_paragraph(
    'Система Anti-Plagiarism — это высокопроизводительное ядро обнаружения плагиата, '
    'написанное на C++17. Основана на шинглинг-алгоритме (shingle-based fingerprinting) '
    'с использованием SimHash для нечёткого дублирования. Система поддерживает '
    'мультитенантность (изоляция по organization_id), многоуровневую индексацию (L1–L5), '
    'и масштабируется горизонтально через шардирование L5 индексов.'
)

doc.add_paragraph(
    'Ключевые возможности:'
)

capabilities = [
    'Индексация текстов на русском, казахском, турецком и английском языках',
    'Нормализация Unicode (lowercase, удаление пунктуации, схлопывание пробелов)',
    'K-шинглинг (K=9) — перекрывающиеся N-граммы токенов',
    'SimHash 128-бит для нечёткого поиска дубликатов',
    'Двоичный формат индекса с O(log N) поиском через binary search',
    'LRU кэширование сегментов в RAM с пин/анпин',
    'Асинхронная очередь задач через PostgreSQL',
    'Компактификация (merge N сегментов → 1)',
    'REST API (HTTP) для индексации и поиска',
    'Docker-контейнеризация с отдельными DEV/PROD средами',
    'Лемматизация казахского языка через Apertium FST (микросервис)',
]

for cap in capabilities:
    p = doc.add_paragraph(cap, style='List Bullet')

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 2. ВЫСОКОУРОВНЕВАЯ АРХИТЕКТУРА
# ═══════════════════════════════════════════════════

doc.add_heading('2. Высокоуровневая архитектура', level=1)

doc.add_paragraph(
    'Система состоит из следующих высокоуровневых компонентов:'
)

arch_text = """┌───────────────────────────────┐
│          Client              │
│  Postman / scripts / UI      │
└─────────────┬────────────────┘
              │ HTTP (REST API)
              ▼
┌────────────────────────────────────────────────┐
│            l5_service (C++ API Server)         │
│  /v1/process  /v1/orgs/.../ingest_zip          │
│  /v1/jobs     /v1/admin/*     /v1/hot_cache/*  │
└──────┬────────────────────────────┬────────────┘
       │ enqueue jobs (PG)         │ read/write index
       ▼                           ▼
┌──────────────┐         ┌──────────────────────┐
│  PostgreSQL  │         │     DATA_ROOT        │
│  job_queue   │         │ orgs/<org>/index/     │
│  documents   │         │   l1/ l2/ l3/ l4/ l5/│
│  hot_cache_* │         │ orgs/<org>/jobs/      │
└──────┬───────┘         └──────────────────────┘
       │ claim_next()
       ▼
┌──────────────────────┐
│  Worker processes    │
│  fork+exec → job     │
│  [Ingest 1..N]       │
│  [Compactor]         │
│  [Admin]             │
└──────────────────────┘"""

add_code_block(doc, arch_text)

doc.add_paragraph()
doc.add_paragraph(
    'Модель исполнения: API-сервер принимает HTTP-запросы и ставит задачи в PostgreSQL. '
    'Worker-процессы непрерывно опрашивают очередь (claim_next), забирают задачу, '
    'делают fork()+exec() себя же в режиме --mode=job и выполняют одну задачу. '
    'После завершения job-процесс завершается — память возвращается ОС.'
)

doc.add_heading('Потоки данных', level=2)

flow_table = [
    ['Клиент → API', 'HTTP POST с текстом / ZIP-архивом', '/v1/process, /v1/orgs/.../ingest_zip'],
    ['API → Postgres', 'Enqueue задачи', 'job_queue таблица'],
    ['Worker → Postgres', 'Claim задачи', 'claim_next() с типом worker'],
    ['Worker → DATA_ROOT', 'Индексация', 'Запись index_native.bin + docids.json'],
    ['API → DATA_ROOT', 'Поиск', 'Чтение сегментов через кэш'],
    ['API → Клиент', 'Результат поиска', 'JSON с hits[], spans[], scores'],
]
add_table(doc, ['Поток', 'Описание', 'Детали'], flow_table, [4, 6, 6])

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 3. СТРУКТУРА ПРОЕКТА
# ═══════════════════════════════════════════════════

doc.add_heading('3. Структура проекта (дерево файлов)', level=1)

tree = """anti-plagiarism/
├── Back_Last/                    ← C++ ядро + HTTP сервис
│   ├── CMakeLists.txt            ← Система сборки CMake
│   ├── Dockerfile                ← Docker-образ
│   ├── docker-compose.yml        ← Оркестрация DEV+PROD
│   ├── entrypoint.sh             ← Точка входа контейнера
│   ├── cpp/
│   │   ├── 1_Text/               ← Нормализация текста
│   │   │   ├── text_common.h/cpp ← Шинглирование, SimHash
│   │   │   ├── extractor.h/cpp   ← Извлечение текста из файлов
│   │   │   └── udpipe_lemma.h    ← UDPipe лемматизация
│   │   ├── 2_Builder/            ← Построение сегментов
│   │   │   └── builder.h/cpp     ← JSONL → сегмент
│   │   ├── 3_Index_IO/           ← Формат, чтение, запись
│   │   │   ├── format.h/cpp      ← HeaderV2, DocMeta, Posting9
│   │   │   ├── manifest.h/cpp    ← Реестр сегментов
│   │   │   ├── reader.h/cpp      ← Загрузка с диска
│   │   │   ├── compactor.h/cpp   ← Компактификация
│   │   │   └── merge.h/cpp       ← K-way merge
│   │   ├── 4_Search/             ← Поиск плагиата
│   │   │   ├── query.h/cpp       ← Построение запроса
│   │   │   ├── search_segment.h  ← Поиск в одном сегменте
│   │   │   ├── search_multi.h    ← Поиск по всем сегментам
│   │   │   ├── result.h/cpp      ← Структуры результата
│   │   │   ├── segobj_cache.h    ← LRU кэш сегментов
│   │   │   └── docids_cache.h    ← LRU кэш docids
│   │   └── 5_Service/            ← HTTP сервис
│   │       ├── main.cpp          ← Точка входа (api/worker/job)
│   │       ├── service.h         ← Оркестратор L5Service
│   │       ├── storage.h/cpp     ← PostgreSQL документы
│   │       ├── tombstone.h/cpp   ← Soft-delete
│   │       ├── core/             ← Ядро сервиса
│   │       │   ├── job_queue.h   ← Очередь задач (PG)
│   │       │   ├── file_lock.h   ← RAII flock()
│   │       │   ├── hot_cache.h   ← mmap/mlock пининг
│   │       │   └── service_internal.h
│   │       └── routes/           ← HTTP маршруты
│   │           ├── route_process.cpp
│   │           ├── route_ingest_zip.cpp
│   │           ├── route_admin_compact.cpp
│   │           └── ...
│   ├── DATA_ROOT_DEV/            ← Данные разработки
│   ├── DATA_ROOT_PROD/           ← Промышленные данные
│   └── initdb/                   ← SQL инициализация
├── archive_test/                 ← Тестовые корпуса и скрипты
│   ├── make_cc100_volumes_from_xz.py  ← Подготовка CC-100
│   ├── tester/
│   │   ├── ingest.py             ← Массовый инджест
│   │   └── ingest_distribute.py  ← Мультитенантный инджест
│   └── cc100_*_tomes/            ← Корпуса (ENG, RU, TR)
├── docs/                         ← PlantUML диаграммы
│   ├── 0_overview.puml
│   ├── 1_indexing.puml
│   ├── 2_storage.puml
│   ├── 3_cache.puml
│   └── 4_search.puml
└── lemma/                        ← Лемматизация
    └── apertium_kaz_fst/         ← Казахский FST анализатор
        ├── api/server.py         ← FastAPI сервис
        ├── tools/build_all.sh    ← Сборка HFST
        └── Dockerfile"""

add_code_block(doc, tree)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 4. МОДУЛЬ 1: TEXT
# ═══════════════════════════════════════════════════

doc.add_heading('4. Модуль 1: Text — Нормализация и токенизация', level=1)

doc.add_paragraph(
    'Модуль cpp/1_Text отвечает за первичную обработку текста: '
    'декодирование, нормализацию Unicode, токенизацию, хэширование токенов и вычисление '
    'шинглов (fingerprints). Это фундамент всей системы.'
)

doc.add_heading('4.1 Нормализация текста', level=2)

doc.add_paragraph(
    'Функция normalize_for_shingles_simple() приводит текст к каноническому виду:'
)

norm_rules = [
    ['1', 'Lowercase ASCII', 'A-Z → a-z'],
    ['2', 'Lowercase кириллица', 'RU/KZ (0x0400..0x052F) → lowercase'],
    ['3', 'Lowercase турецкий', 'ç, ğ, ı, ö, ş, ü → lowercase'],
    ['4', 'Арабские цифры', '٠-٩ → 0-9'],
    ['5', 'Пунктуация → пробел', 'Все знаки препинания заменяются на пробел'],
    ['6', 'Схлопывание пробелов', 'Множественные пробелы → один'],
    ['7', 'Обрезка', 'Удаление завершающего пробела'],
]
add_table(doc, ['#', 'Правило', 'Описание'], norm_rules, [1.5, 5, 7])

doc.add_heading('4.2 Токенизация', level=2)

doc.add_paragraph(
    'Функция tokenize_spans() разбивает нормализованный текст на токены по пробельным символам '
    '(ASCII пробел + NBSP). Возвращает вектор TokenSpan{start, len} — позиции токенов в тексте.'
)

doc.add_heading('4.3 Хэширование', level=2)

doc.add_paragraph(
    'Каждый отдельный токен хэшируется через XXHash64 (функция hash_tokens_bytes_spans). '
    'Это позволяет избежать повторного хэширования одних и тех же байтов при построении шинглов.'
)

doc.add_heading('4.4 Шинглирование (K=9)', level=2)

doc.add_paragraph(
    'Шингл — это хэш K=9 последовательных токенов. Для позиции pos от 0 до (N_tokens - 9) '
    'вычисляется hash_shingle_token_hashes(token_hashes[pos..pos+9]). '
    'Каждый шингл порождает запись Posting9{hash, doc_id, position}.'
)

doc.add_paragraph(
    'Ограничения: max_tokens_per_doc = 100 000, max_shingles_per_doc = 50 000, '
    'shingle_stride = 1 (по умолчанию).'
)

doc.add_heading('4.5 SimHash 128-бит', level=2)

doc.add_paragraph(
    'SimHash — это Locality-Sensitive Hash (LSH). Для каждого документа вычисляется '
    '128-битный fingerprint (hi: uint64, lo: uint64) по всем хэшам токенов. '
    'Используется для быстрого обнаружения near-duplicates: '
    'документы с малым расстоянием Хэмминга между SimHash являются кандидатами.'
)

doc.add_heading('4.6 Извлечение текста (extractor)', level=2)

doc.add_paragraph(
    'Extractor читает текстовые файлы (.txt) с диска. '
    'Поддерживает UTF-8 и fallback CP1251→UTF-8. '
    'Жёсткий лимит: 256 MiB на файл. '
    'Ключевое правило: extractor НИКОГДА не нормализует текст — '
    'он только декодирует и обрезает.'
)

doc.add_heading('4.7 UDPipe лемматизация (опционально)', level=2)

doc.add_paragraph(
    'При включённой опции L5_ENABLE_UDPIPE модуль загружает модели UDPipe '
    'для русского, казахского, английского и турецкого языков. '
    'Функция udlang_guess_from_norm() определяет язык эвристически. '
    'Лемматизация выполняется через класс UdpipeRuntime (singleton).'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 5. МОДУЛЬ 2: BUILDER
# ═══════════════════════════════════════════════════

doc.add_heading('5. Модуль 2: Builder — Построение сегментов', level=1)

doc.add_paragraph(
    'Модуль cpp/2_Builder строит сегменты инвертированного индекса из корпуса JSONL. '
    'Основная функция — build_segment_jsonl().'
)

doc.add_heading('5.1 Параметры сборки (BuildOptions)', level=2)

build_opts = [
    ['segment_name', 'Имя сегмента', 'auto-generated'],
    ['max_text_bytes_per_doc', 'Макс. размер текста', '8 MiB'],
    ['max_tokens_per_doc', 'Макс. токенов', '100 000'],
    ['max_shingles_per_doc', 'Макс. шинглов', '50 000'],
    ['max_docs_in_segment', 'Макс. документов в сегменте', 'без лимита'],
    ['shingle_stride', 'Шаг шинглирования', '1'],
    ['max_threads', 'Потоки', '16'],
    ['ram_limit_bytes', 'Лимит RAM', '512 MiB'],
    ['build_lsh', 'Строить LSH индекс', 'true'],
    ['lsh_blocks', 'Блоки LSH', '8'],
]
add_table(doc, ['Параметр', 'Описание', 'По умолчанию'], build_opts, [5, 6, 3])

doc.add_heading('5.2 Процесс построения', level=2)

doc.add_paragraph(
    'Пайплайн построения сегмента (streaming, ограниченная память):'
)

steps = [
    'Чтение JSONL файла (corpus) построчно',
    'Нормализация текста (normalize_for_shingles)',
    'Токенизация (tokenize_spans)',
    'Хэширование токенов (hash_tokens_bytes_spans)',
    'Построение шинглов K=9 (hash_shingle_token_hashes)',
    'Вычисление SimHash128 (simhash128_token_hashes)',
    'Создание Posting9{h, did, pos} для каждого шингла',
    'Создание DocMeta{tok_len, simhash_hi, simhash_lo}',
    'Сортировка постингов по (h, did, pos)',
    'Запись index_native.bin (Header + DocMeta[] + Posting9[])',
    'Запись index_native_docids.json (DocInfo[])',
    'Обновление manifest.jsonl',
]

for i, step in enumerate(steps, 1):
    p = doc.add_paragraph(f'{i}. {step}')

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 6. МОДУЛЬ 3: INDEX_IO
# ═══════════════════════════════════════════════════

doc.add_heading('6. Модуль 3: Index_IO — Формат хранения на диске', level=1)

doc.add_paragraph(
    'Модуль cpp/3_Index_IO определяет бинарный формат индексных файлов, '
    'реализует чтение, запись, валидацию, merge и компактификацию сегментов.'
)

doc.add_heading('6.1 Файлы сегмента', level=2)

seg_files = [
    ['index_native.bin', 'Основной бинарный индекс', 'Header + DocMeta[] + Posting9[]'],
    ['index_native_docids.json', 'Метаданные документов', 'JSON массив DocInfo'],
    ['manifest.jsonl', 'Реестр всех сегментов', 'По одной строке на сегмент'],
]
add_table(doc, ['Файл', 'Назначение', 'Содержимое'], seg_files, [5, 5, 6])

doc.add_heading('6.2 Бинарный формат index_native.bin', level=2)

doc.add_paragraph('Файл состоит из трёх последовательных блоков:')

format_text = """╔═══════════════════════════════════╗
║ HeaderV2 (28 байт)                ║
╠═══════════════════════════════════╣
║ magic[4] = "PLAG"                 ║
║ version   = 2                     ║
║ n_docs    : uint32                ║
║ n_post9   : uint64                ║
║ n_post13  : uint64 (зарезерв.=0) ║
╠═══════════════════════════════════╣
║ DocMeta[n_docs] (20 байт каждый)  ║
║   tok_len    : uint32             ║
║   simhash_hi : uint64             ║
║   simhash_lo : uint64             ║
╠═══════════════════════════════════╣
║ Posting9[n_post9] (16 байт каждый)║
║   h   : uint64 (хэш шингла)      ║
║   did : uint32 (ID документа)     ║
║   pos : uint32 (позиция токена)   ║
╚═══════════════════════════════════╝"""

add_code_block(doc, format_text)

doc.add_paragraph(
    'Размер файла = 28 + (20 × N_docs) + (16 × N_shingles) байт. '
    'Пример: 100K документов, 50M шинглов → 28 + 2 MB + 800 MB ≈ 802 MB.'
)

doc.add_paragraph(
    'Постинги отсортированы по (h, did, pos) для бинарного поиска O(log P).'
)

doc.add_heading('6.3 DocInfo (JSON)', level=2)

docinfo_fields = [
    ['doc_id', 'Внутренний уникальный ID'],
    ['module', 'Модуль (scope для фильтрации)'],
    ['organization_id', 'Тенант (изоляция)'],
    ['external_id', 'Внешний идентификатор'],
    ['source_path', 'Путь к оригиналу'],
    ['source_name', 'Имя оригинального файла'],
    ['preview_text', 'Первые 240 символов документа'],
]
add_table(doc, ['Поле', 'Описание'], docinfo_fields, [5, 10])

doc.add_heading('6.4 Manifest', level=2)

doc.add_paragraph(
    'Файл manifest.jsonl — реестр сегментов. Каждая строка:'
)

add_code_block(doc, '{"segment_name": "seg_xxx", "path": "seg_xxx/",\n "built_at_utc": "20260209_120000",\n "stats": {"docs": 1000, "k9": 50000}}')

doc.add_heading('6.5 Компактификация', level=2)

doc.add_paragraph(
    'Компактификация (compactor.cpp) объединяет N мелких сегментов в один крупный. '
    'Алгоритм: K-way merge отсортированных постингов без повторной токенизации. '
    'Параметр fanout (по умолчанию 20) определяет, сколько сегментов сливать за раз. '
    'После слияния старые сегменты удаляются, манифест обновляется атомарно.'
)

doc.add_heading('6.6 Валидация', level=2)

doc.add_paragraph(
    'Функция header_v2_sane() проверяет заголовок индекса против фактического размера файла '
    'для предотвращения OOM на повреждённых данных.'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 7. МОДУЛЬ 4: SEARCH
# ═══════════════════════════════════════════════════

doc.add_heading('7. Модуль 4: Search — Поиск плагиата', level=1)

doc.add_paragraph(
    'Модуль cpp/4_Search реализует полный пайплайн поиска: '
    'от построения запроса до выдачи ранжированных результатов.'
)

doc.add_heading('7.1 Построение запроса', level=2)

doc.add_paragraph(
    'Функция build_query_shingles() обрабатывает текст запроса тем же способом, '
    'что и при индексации: нормализация → токенизация → хэширование → шинглирование K=9. '
    'Результат — QueryShingles{items[], total_shingles, tok_len, simhash}.'
)

doc.add_paragraph(
    'Одинаковые хэши группируются в QueryHash{h, vector<qpos>} — '
    'один хэш может встретиться в нескольких позициях запроса.'
)

doc.add_heading('7.2 Поиск в сегменте', level=2)

doc.add_paragraph(
    'Функция search_in_segment() выполняет следующие шаги:'
)

search_steps = [
    'Binary search: для каждого уникального хэша запроса (U штук) → range_for_hash_safe() '
    'находит диапазон [lower_bound, upper_bound) в отсортированном массиве постингов. O(log P).',
    'Сбор совпадений: для каждого постинга в найденном диапазоне записывается (did, pos). O(M).',
    'SearchScratch (Thread-Local): epoch-stamped массивы hits[did]. '
    'Позволяет избежать O(N_docs) обнуления — реальная стоимость O(M).',
    'Span Detection: для каждого документа с совпадениями собираются все пары (qpos, dpos), '
    'сортируются по delta = dpos - qpos. Одинаковый delta = диагональное совпадение.',
    'Группировка по delta → построение SpanTmp{q_start, q_end, d_start, d_end, len_shingles}. '
    'Смежные позиции объединяются (gap allowed: span_gap).',
    'Scoring: union_len_inclusive_from_spans() вычисляет покрытие. '
    'overlap_ratio = C / union_len. Формируется Hit.',
]
for i, s in enumerate(search_steps, 1):
    p = doc.add_paragraph(f'{i}. {s}')

doc.add_heading('7.3 Параметры поиска (SearchOptions)', level=2)

search_opts = [
    ['topk', 'Максимум результатов', '20'],
    ['candidates_topn', 'Кандидатов для детального анализа', '200'],
    ['min_hits', 'Минимум совпадений для документа', '2'],
    ['max_postings_per_hash', 'Макс. постингов на один хэш', '50 000'],
    ['span_min_len', 'Мин. длина span (в шинглах)', '6'],
    ['span_gap', 'Допустимый разрыв в span', '0'],
    ['max_spans_per_doc', 'Макс. spans на документ', '10'],
    ['alpha', 'Порог совпадения', '0.60'],
]
add_table(doc, ['Параметр', 'Описание', 'По умолчанию'], search_opts, [5, 6, 3])

doc.add_heading('7.4 Мульти-сегментный поиск', level=2)

doc.add_paragraph(
    'Функция search_out_root() загружает manifest, итерирует все сегменты, '
    'вызывает search_in_segment() для каждого и объединяет результаты. '
    'Одинаковый doc_id из разных сегментов → берётся лучший C (coverage). '
    'Финальная сортировка по C по убыванию → top-K.'
)

doc.add_heading('7.5 Структура результата', level=2)

result_fields = [
    ['Hit.doc_id', 'ID документа'],
    ['Hit.C', 'Общий процент совпадения'],
    ['Hit.Cq', 'Покрытие запроса (%)'],
    ['Hit.Cd', 'Покрытие документа (%)'],
    ['Hit.match_spans[]', 'Массив совпавших интервалов'],
    ['MatchSpan.q_from/q_to', 'Позиции в запросе (токены)'],
    ['MatchSpan.d_from/d_to', 'Позиции в документе (токены)'],
    ['MatchSpan.length', 'Длина совпадения'],
    ['SearchResult.segments_scanned', 'Количество просканированных сегментов'],
]
add_table(doc, ['Поле', 'Описание'], result_fields, [6, 10])

doc.add_heading('7.6 Вычислительная сложность', level=2)

complexity = [
    ['Обработка запроса', 'O(Q_tokens + Q_shingles)'],
    ['Один сегмент', 'O(U × log P + M)'],
    ['Span Detection', 'O(M × log M)'],
    ['Общая', 'O(S × (U × log P + M))'],
]
add_table(doc, ['Этап', 'Сложность'], complexity, [6, 10])

doc.add_paragraph(
    'Где: S = сегменты, U = уникальные хэши запроса, P = постинги в сегменте, M = найденные совпадения.'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 8. МОДУЛЬ 5: SERVICE
# ═══════════════════════════════════════════════════

doc.add_heading('8. Модуль 5: Service — HTTP-сервис и воркеры', level=1)

doc.add_paragraph(
    'Модуль cpp/5_Service — верхний уровень системы. Содержит HTTP-сервер, '
    'систему воркеров, маршрутизацию запросов и оркестрацию.'
)

doc.add_heading('8.1 Режимы работы', level=2)

modes = [
    ['--mode=api', 'HTTP сервер', 'Слушает запросы (0.0.0.0:8088), регистрирует роуты, '
     'управляет hot cache. На старте применяет сохранённые пины.'],
    ['--mode=worker', 'Воркер', 'Циклически опрашивает job_queue через claim_next(). '
     'Делает fork()+exec() себя в режиме --mode=job. Типы: all, ingest, compact, admin.'],
    ['--mode=job', 'Одиночная задача', 'Выполняет одну задачу по ID из БД. '
     'Типы: FLUSH_L1, L1_BUILD, INDEX_TEXT, INGEST_L5_ZIP, COMPACT_SMALL, COMPACT_L5, '
     'REBUILD_ONE, WIPE_ALL.'],
]
add_table(doc, ['Режим', 'Роль', 'Описание'], modes, [3, 3, 10])

doc.add_heading('8.2 Класс L5Service', level=2)

doc.add_paragraph(
    'L5Service — главный оркестратор, объединяющий все модули:'
)

svc_methods = [
    ['index_text_document()', 'Индексация одного текста через API'],
    ['ingest_l5_zip_build_segment()', 'Извлечение ZIP → сборка L5 сегмента'],
    ['ingest_l5_fs_dirs_build_segment()', 'Инджест из файловой системы'],
    ['search_levels()', 'Поиск по уровням L1–L4 и L5 (с шардами)'],
    ['compact_small_levels()', 'Компактификация мелких сегментов (L1–L4)'],
    ['compact_l5_shards()', 'Компактификация L5 шардов'],
    ['flush_l1_tail_job()', 'Запечатывание L1 буфера при простое'],
    ['l1_build_segment_job()', 'Сборка L1 сегмента из JSONL'],
]
add_table(doc, ['Метод', 'Описание'], svc_methods, [7, 9])

doc.add_heading('8.3 Storage (PostgreSQL)', level=2)

doc.add_paragraph(
    'Класс Storage — обёртка над libpq для работы с таблицей documents. '
    'Методы: upsert_doc_get_id (upsert с возвратом internal_id), '
    'get_by_source_id, list_docs, mark_deleted_by_source_id, '
    'update_last_segment_by_ids и др.'
)

doc.add_heading('8.4 Tombstones', level=2)

doc.add_paragraph(
    'Soft-delete реализован через файл tombstones.jsonl per-org. '
    'Класс Tombstones: load() → unordered_set<string>, append() добавляет, '
    'contains() проверяет. Удалённые документы пропускаются при поиске.'
)

doc.add_heading('8.5 Блокировки', level=2)

doc.add_paragraph(
    'FileLock — RAII обёртка над flock(). Два уровня: '
    'глобальный (.global.lock — Shared для обычных jobs, Exclusive для WIPE_ALL) '
    'и per-org (эксклюзивный на индексацию/компактификацию).'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 9. КЭШИРОВАНИЕ
# ═══════════════════════════════════════════════════

doc.add_heading('9. Система кэширования (Hot Cache)', level=1)

doc.add_heading('9.1 SegmentObjectCache (RAM)', level=2)

doc.add_paragraph(
    'LRU кэш для загруженных сегментов (SegmentData = header + DocMeta[] + Posting9[]). '
    'Конфигурация: PLAGIO_SEGOBJ_CACHE_BYTES (env). По умолчанию: без лимита. '
    'Рекомендация: 60 GiB для крупных индексов.'
)

cache_features = [
    'LRU eviction (least recently used)',
    'Pin/Unpin для горячих сегментов',
    'Валидация mtime (перезагрузка при изменении)',
    'Thread-safe get_or_load()',
    'Inflight dedup (одна загрузка на сегмент)',
]
for f in cache_features:
    doc.add_paragraph(f, style='List Bullet')

doc.add_paragraph(
    'Память на сегмент: docmeta = 20 × N_docs байт, postings = 16 × N_shingles байт. '
    'Пример: 1M документов, 500M шинглов → ~8 GB на сегмент.'
)

doc.add_heading('9.2 DocidsCache (RAM)', level=2)

doc.add_paragraph(
    'LRU кэш для массивов DocInfo (строки). Аналогичная архитектура. '
    'PLAGIO_DOCIDS_CACHE_BYTES (env). Расход: ~300–500 байт/документ.'
)

doc.add_heading('9.3 HotCache (OS Page Cache)', level=2)

doc.add_paragraph(
    'Класс HotCache использует mmap/mlock/read для закрепления файлов '
    'в страничном кэше ОС. Методы: pin_file(), pin_segment_dir() (bin/docids/meta), '
    'unpin_file(), refresh_best_effort(). '
    'Состояние пинов сохраняется в PostgreSQL (таблицы hot_cache_pins, hot_cache_events).'
)

doc.add_heading('9.4 Поток кэширования', level=2)

cache_flow = """Search Request
    │
    ▼
get_or_load(segment_path)
    │
    ├── Cache HIT → shared_ptr<SegmentData>
    │
    └── Cache MISS
         │
         ▼
    Load from SSD (reader)
         │
         ▼
    Check budget → LRU Evict (skip pinned)
         │
         ▼
    Add to cache → return shared_ptr"""

add_code_block(doc, cache_flow)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 10. POSTGRESQL
# ═══════════════════════════════════════════════════

doc.add_heading('10. PostgreSQL — метаданные и очередь задач', level=1)

doc.add_heading('10.1 Таблица job_queue', level=2)

doc.add_paragraph(
    'Очередь асинхронных задач. Воркеры забирают задачи через claim_next() '
    'с фильтрацией по типу (ingest/compact/admin).'
)

job_fields = [
    ['id', 'SERIAL PRIMARY KEY'],
    ['type', 'VARCHAR — тип задачи (INDEX_TEXT, COMPACT_SMALL, ...)'],
    ['status', 'VARCHAR — pending/running/done/failed'],
    ['payload', 'JSONB — входные параметры'],
    ['result', 'JSONB — результат выполнения'],
    ['error', 'TEXT — сообщение об ошибке'],
    ['created_at', 'TIMESTAMP'],
    ['updated_at', 'TIMESTAMP'],
]
add_table(doc, ['Поле', 'Описание'], job_fields, [4, 12])

doc.add_heading('10.2 Таблица documents', level=2)

doc.add_paragraph(
    'Метаданные проиндексированных документов:'
)

doc_fields = [
    ['id', 'SERIAL PRIMARY KEY (internal ID)'],
    ['org_id', 'ID организации'],
    ['source_id', 'Внешний ID документа'],
    ['file_name', 'Имя файла'],
    ['title', 'Заголовок'],
    ['author', 'Автор'],
    ['deleted', 'BOOLEAN — soft-delete'],
    ['last_segment', 'Последний сегмент, содержащий документ'],
    ['current_segment', 'Текущий сегмент'],
    ['module', 'Модуль (scope)'],
]
add_table(doc, ['Поле', 'Описание'], doc_fields, [4, 12])

doc.add_heading('10.3 Таблицы hot cache (опционально)', level=2)

doc.add_paragraph(
    'hot_cache_events — история событий пин/анпин. '
    'hot_cache_pins — текущие пины. '
    'hot_cache_policy — политики кэширования.'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 11. ПАЙПЛАЙН ИНДЕКСАЦИИ
# ═══════════════════════════════════════════════════

doc.add_heading('11. Пайплайн индексации (Indexing Pipeline)', level=1)

doc.add_paragraph(
    'Полный путь текста от загрузки до записи на диск:'
)

pipeline_text = """┌─────────────┐
│  Raw Text   │  UTF-8, max 8 MiB, из JSONL корпуса
└──────┬──────┘
       │
       ▼
┌──────────────────────────────┐
│ normalize_for_shingles()     │  lowercase + пунктуация → пробел
└──────┬───────────────────────┘
       │
       ▼
┌──────────────────────────────┐
│ tokenize_spans()             │  Разбивка по пробелам → TokenSpan[]
└──────┬───────────────────────┘
       │
       ▼
┌──────────────────────────────┐
│ hash_tokens_bytes_spans()    │  XXHash64 каждого токена
└──────┬───────────────────────┘
       │
       ├──────────────┐
       ▼              ▼
┌─────────────┐ ┌──────────────┐
│ Shingles    │ │ SimHash128   │
│ K=9, stride │ │ fingerprint  │
│ → Posting9  │ │ → DocMeta    │
└──────┬──────┘ └──────┬───────┘
       │               │
       ▼               ▼
┌──────────────────────────────┐
│ Sort postings (h, did, pos)  │
└──────┬───────────────────────┘
       │
       ▼
┌──────────────────────────────┐
│ Write segment files:         │
│  index_native.bin            │
│  index_native_docids.json    │
│  manifest.jsonl (append)     │
└──────────────────────────────┘"""

add_code_block(doc, pipeline_text)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 12. ПАЙПЛАЙН ПОИСКА
# ═══════════════════════════════════════════════════

doc.add_heading('12. Пайплайн поиска (Search Pipeline)', level=1)

search_pipeline = """┌─────────────┐
│ Query Text  │  max 8 MiB
└──────┬──────┘
       │
       ▼
┌──────────────────────────────┐
│ build_query_shingles()       │
│  normalize → tokenize →      │
│  hash → shingles K=9         │
│  → QueryShingles{items[]}    │
└──────┬───────────────────────┘
       │
       ▼
┌──────────────────────────────┐
│ load_manifest(out_root)      │  Получить список сегментов
└──────┬───────────────────────┘
       │
       ▼  Для каждого сегмента:
┌──────────────────────────────┐
│ segobj_cache.get_or_load()   │  Загрузить postings в RAM
│ docids_cache.get_or_load()   │  Загрузить DocInfo в RAM
└──────┬───────────────────────┘
       │
       ▼
┌──────────────────────────────┐
│ search_in_segment()          │
│  1. Binary search по хэшам   │  O(U × log P)
│  2. Сбор совпадений          │  O(M)
│  3. Span Detection           │  Группировка по delta
│  4. Scoring (C, Cq, Cd)      │  overlap_ratio
└──────┬───────────────────────┘
       │
       ▼
┌──────────────────────────────┐
│ Merge hits across segments   │  Лучший C для doc_id
│ Sort by C descending         │
│ Return top-K                 │
└──────────────────────────────┘
       │
       ▼
┌──────────────────────────────┐
│ SearchResult {               │
│   query, segments_scanned,   │
│   hits[]: {doc_id, C, Cq,    │
│     Cd, spans[], metadata}   │
│ }                            │
└──────────────────────────────┘"""

add_code_block(doc, search_pipeline)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 13. ФОРМАТ ХРАНЕНИЯ
# ═══════════════════════════════════════════════════

doc.add_heading('13. Формат хранения данных на диске', level=1)

doc.add_heading('13.1 Иерархия DATA_ROOT', level=2)

data_root = """DATA_ROOT/
├── orgs/
│   └── <org_id>/
│       ├── index/
│       │   ├── l1/              ← Уровень 1 (per-document)
│       │   │   ├── seg_xxx/
│       │   │   │   ├── index_native.bin
│       │   │   │   └── index_native_docids.json
│       │   │   └── manifest.jsonl
│       │   ├── l2/              ← Уровень 2
│       │   ├── l3/              ← Уровень 3
│       │   ├── l4/              ← Уровень 4
│       │   └── l5/              ← Уровень 5 (bulk sharded)
│       │       ├── s00/         ← Шард 0
│       │       ├── s01/         ← Шард 1
│       │       └── sNN/
│       ├── jobs/                ← Входные данные tasks
│       │   └── job_<id>/input.txt
│       ├── uploads/             ← Оригиналы (debug)
│       └── tombstones.jsonl     ← Soft-delete
└── .global.lock                 ← Глобальная блокировка"""

add_code_block(doc, data_root)

doc.add_heading('13.2 Уровни индексации', level=2)

levels = [
    ['L1', 'Per-document', 'Каждый текст → отдельный маленький сегмент. Быстрая индексация.'],
    ['L2', 'Compacted L1', 'Результат компактификации нескольких L1 сегментов.'],
    ['L3', 'Compacted L2', 'Дальнейшее укрупнение сегментов.'],
    ['L4', 'Compacted L3', 'Ещё более крупные сегменты.'],
    ['L5', 'Bulk ingest', 'Массовая загрузка корпусов (ZIP/filesystem). Поддержка шардирования.'],
]
add_table(doc, ['Уровень', 'Тип', 'Описание'], levels, [2, 4, 10])

doc.add_heading('13.3 Константы формата', level=2)

consts = [
    ['K_SHINGLE', '9', 'Ширина шингла (в токенах)'],
    ['HEADER_V2_BYTES', '28', 'Размер заголовка файла'],
    ['DOCMETA_V2_BYTES', '20', 'Размер записи DocMeta'],
    ['POSTING9_V2_BYTES', '16', 'Размер записи Posting9'],
    ['Magic', '"PLAG"', '4 байта идентификатор файла'],
    ['Version', '2', 'Версия формата'],
]
add_table(doc, ['Константа', 'Значение', 'Описание'], consts, [5, 3, 7])

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 14. ПОДГОТОВКА КОРПУСА
# ═══════════════════════════════════════════════════

doc.add_heading('14. Подготовка корпуса и инджест скрипты', level=1)

doc.add_heading('14.1 Подготовка CC-100 корпуса', level=2)

doc.add_paragraph(
    'Скрипт make_cc100_volumes_from_xz.py конвертирует сжатый (.txt.xz) корпус CC-100 '
    'в структурированное дерево документов:'
)

add_code_block(doc, 'vol_NNNN/CCC/doc_DDDDDD.txt\n\nНастройки по умолчанию:\n  volumes: 15\n  chunks per volume: 100\n  docs per chunk: 1000\n  min chars per doc: 200K\n  → до 1.5M документов')

doc.add_heading('14.2 Массовый инджест (ingest.py)', level=2)

doc.add_paragraph(
    'Скрипт archive_test/tester/ingest.py поддерживает два режима:'
)

ingest_modes = [
    ['--target l5', 'Пакует .txt файлы в ZIP → POST /v1/orgs/.../l5/ingest_zip', 'Шинглинг L5'],
    ['--target l14', 'POST каждого файла на /v1/process с do_index=True', 'Per-document L1-L4'],
]
add_table(doc, ['Режим', 'Механизм', 'Индекс'], ingest_modes, [3, 8, 4])

doc.add_paragraph(
    'Особенности: JSONL-idempotency (повторный запуск пропускает обработанные), '
    'retry с backoff, настраиваемый concurrency.'
)

doc.add_heading('14.3 Мультитенантный инджест (ingest_distribute.py)', level=2)

doc.add_paragraph(
    'Распределяет CC-100 RU корпус по 4 организациям (тенантам) для тестирования '
    'мультитенантности. Chunk 000 → org "1", 001 → org "2", и т.д.'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 15. ЛЕММАТИЗАЦИЯ
# ═══════════════════════════════════════════════════

doc.add_heading('15. Лемматизация — Apertium Kazakh FST', level=1)

doc.add_paragraph(
    'Отдельный микросервис для морфологического анализа казахского языка, '
    'построенный на Apertium + HFST (Helsinki Finite-State Technology).'
)

doc.add_heading('15.1 Архитектура', level=2)

lemma_arch = [
    ['Компонент', 'Технология', 'Описание'],
]
lemma_rows = [
    ['Морфоанализатор', 'HFST / Apertium', 'FST transducer kaz.automorf.hfst'],
    ['Генератор', 'HFST / Apertium', 'FST transducer kaz.autogen.hfst'],
    ['API', 'FastAPI (Python)', 'HTTP эндпоинты /analyze, /spell-check'],
    ['Кэширование', 'ai_cache.py', 'Кэш результатов анализа'],
    ['Docker', 'Порт 8000', '8 воркеров, timeout 10s'],
]
add_table(doc, ['Компонент', 'Технология', 'Описание'], lemma_rows, [4, 4, 8])

doc.add_heading('15.2 Производительность', level=2)

doc.add_paragraph(
    'Чистый FST (без ML): ~50 000+ слов/сек, <1ms на слово. '
    'Нагрузочное тестирование: 100 потоков, 2000 запросов.'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 16. DOCKER
# ═══════════════════════════════════════════════════

doc.add_heading('16. Docker и деплой', level=1)

doc.add_heading('16.1 Docker-образ', level=2)

doc.add_paragraph(
    'Образ основан на ubuntu:24.04. Содержит runtime-библиотеки '
    '(libzip4, libpq5, libsimdjson19, zip) и скомпилированный бинарник l5_service.'
)

doc.add_heading('16.2 Docker Compose: DEV и PROD', level=2)

compose_services = [
    ['postgres_dev', 'PostgreSQL 16', 'Порт 5433, init scripts'],
    ['api_dev', 'API сервер (DEV)', 'Порт 8088, 128 GiB sort RAM, 32 L5 шарда'],
    ['worker_ingest_dev_1..4', '4 инджест-воркера', 'Параллельная индексация'],
    ['worker_compact_dev', 'Компактор', 'Merge мелких сегментов'],
    ['worker_admin_dev', 'Админ воркер', 'WIPE_ALL и admin ops'],
    ['api_prod', 'API сервер (PROD)', 'Порт 8089, БД plagio_prod'],
    ['worker_ingest_prod', '1 инджест-воркер', 'Prod индексация'],
    ['worker_compact_prod', 'Компактор', 'Prod merge'],
    ['worker_admin_prod', 'Админ воркер', 'Prod admin'],
]
add_table(doc, ['Сервис', 'Роль', 'Детали'], compose_services, [5, 4, 7])

doc.add_heading('16.3 Переменные окружения', level=2)

env_vars = [
    ['PLAGIO_PG_DSN', 'DSN подключения к PostgreSQL'],
    ['DATA_ROOT', 'Корневая директория данных'],
    ['PLAGIO_BIND_HOST', 'Адрес привязки (по умолчанию 0.0.0.0)'],
    ['PLAGIO_PORT', 'Порт (по умолчанию 8088)'],
    ['PLAGIO_SEGOBJ_CACHE_BYTES', 'Бюджет RAM-кэша сегментов'],
    ['PLAGIO_DOCIDS_CACHE_BYTES', 'Бюджет RAM-кэша DocInfo'],
    ['PLAGIO_BUILD_THREADS', 'Потоки сборки сегмента'],
    ['PLAGIO_SORT_RAM_BYTES', 'RAM для сортировки постингов'],
    ['PLAGIO_WORKER_PROCS', 'Количество worker-процессов в контейнере'],
]
add_table(doc, ['Переменная', 'Описание'], env_vars, [6, 10])

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 17. API РОУТЫ
# ═══════════════════════════════════════════════════

doc.add_heading('17. API роуты', level=1)

routes = [
    ['POST', '/v1/process', 'Индексация + поиск', 'Основной роут: do_index/do_search'],
    ['POST', '/v1/orgs/{org}/l5/ingest_zip', 'Инджест ZIP', 'multipart upload → L5 сегмент'],
    ['GET', '/v1/jobs/{id}', 'Статус задачи', 'payload/result/error из PG'],
    ['POST', '/v1/orgs/{org}/admin/compact_small', 'Компактификация', 'fanout=N'],
    ['POST', '/v1/orgs/{org}/admin/compact_l5', 'Компактификация L5', 'fanout=N'],
    ['POST', '/v1/orgs/{org}/l5/rebuild_one', 'Rebuild', 'Консолидация шардов'],
    ['GET', '/v1/orgs/{org}/l5/rebuild_info', 'Rebuild info', 'rebuild_last.json'],
    ['GET', '/v1/segment_docs', 'Список документов', 'Скан индексов на диске'],
    ['GET', '/v1/orgs/{org}/debug/normalized_text', 'Debug', 'Нормализованный текст'],
    ['GET', '/v1/orgs/{org}/admin/hot_cache/status', 'Кэш статус', 'Состояние RAM кэша'],
    ['POST', '/v1/orgs/{org}/admin/hot_cache/pin', 'Пин кэша', 'Закрепить сегменты в RAM'],
    ['POST', '/v1/orgs/{org}/admin/hot_cache/clear', 'Очистка кэша', 'file|segobj|both'],
    ['POST', '/v1/admin/wipe_all', 'Полная очистка', 'confirm=WIPE_ALL'],
]
add_table(doc, ['Метод', 'Роут', 'Назначение', 'Описание'], routes, [2, 6, 3.5, 5])

doc.add_heading('17.1 Пример запроса: индексация', level=2)

add_code_block(doc, '''POST /v1/process HTTP/1.1
Content-Type: application/json

{
  "organization_id": 28,
  "document_id": "doc_1",
  "file_name": "doc_1.txt",
  "text": "Пример текста для индексации...",
  "do_index": true,
  "do_search": false
}''')

doc.add_heading('17.2 Пример запроса: поиск', level=2)

add_code_block(doc, '''POST /v1/process HTTP/1.1
Content-Type: application/json

{
  "organization_id": 28,
  "document_id": "q1",
  "file_name": "q1.txt",
  "text": "Текст для проверки на плагиат...",
  "do_index": false,
  "do_search": true
}''')

doc.add_page_break()

# ═══════════════════════════════════════════════════
# 18. РЕСУРСЫ И ПРОИЗВОДИТЕЛЬНОСТЬ
# ═══════════════════════════════════════════════════

doc.add_heading('18. Ресурсы и производительность', level=1)

doc.add_heading('18.1 Объём данных на диске', level=2)

disk_data = [
    ['index_native.bin', '28 + 20×D + 16×S байт', 'D=docs, S=shingles'],
    ['docids.json', '~200–500 байт/doc', 'JSON строки'],
    ['Пример: 100K docs, 50M shingles', '~802 MB', 'Один сегмент'],
]
add_table(doc, ['Компонент', 'Размер', 'Примечание'], disk_data, [6, 5, 5])

doc.add_heading('18.2 Потребление RAM', level=2)

ram_data = [
    ['Сегмент в кэше', '20×D + 16×S байт', 'То же что SSD'],
    ['DocInfo в кэше', '~300–500 байт/doc', 'Строки'],
    ['1M docs, 500M shingles', '~8 GB', 'Один сегмент в RAM'],
    ['Рекомендация segobj cache', '60 GiB', 'Для крупных индексов'],
]
add_table(doc, ['Компонент', 'Размер', 'Примечание'], ram_data, [6, 5, 5])

doc.add_heading('18.3 Вычислительная сложность', level=2)

cpu_data = [
    ['Индексация', 'O(text + tokens + shingles)', 'Линейная по тексту'],
    ['Поиск (1 сегмент)', 'O(U × log P + M)', 'U=хэши, P=постинги, M=совпадения'],
    ['Поиск (все)', 'O(S × (U × log P + M))', 'S=сегменты'],
    ['Компактификация', 'O(K × N)', 'K-way merge, N=суммарные постинги'],
]
add_table(doc, ['Операция', 'Сложность', 'Примечание'], cpu_data, [5, 5, 6])

doc.add_heading('18.4 Лимиты', level=2)

limits = [
    ['Макс. текст на документ', '8 MiB'],
    ['Макс. токенов на документ', '100 000'],
    ['Макс. шинглов на документ', '50 000'],
    ['Макс. файл (extractor)', '256 MiB'],
    ['K (ширина шингла)', '9 токенов'],
    ['По умолчанию top-K результатов', '20'],
    ['Fanout компактификации', '20'],
    ['L5 шардов (docker)', '32'],
    ['Worker потоков по умолчанию', '20'],
]
add_table(doc, ['Параметр', 'Значение'], limits, [7, 6])

# ═══════════════════════════════════════════════════
# Сохранение
# ═══════════════════════════════════════════════════

output_path = '/home/oysyn/Date_Index/anti-plagiarism/Anti_Plagiarism_Architecture.docx'
doc.save(output_path)
print(f'Документ сохранён: {output_path}')
